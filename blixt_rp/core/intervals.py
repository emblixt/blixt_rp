"""
Class to handle (working) intervals and tops
It should have general info about an interval (e.g. name, color, ...)
in one class, and then a collection of these intervals in another class
that attaches a well and a depth to each interval
"""
import numpy as np
import sys
import os
import logging
import pint
import pandas as pd
from numpy.distutils.command.egg_info import egg_info
from bokeh.plotting import figure, show

from .. import ureg, Q_

logger = logging.getLogger(__name__)

# Add to path to avoid having to install libraries, useful in development
project_dir = os.path.dirname(__file__).replace('blixt_rp\\blixt_rp\\core','')
sys.path.append(os.path.join(str(project_dir), 'blixt_utils'))

from blixt_utils.misc.attribdict import AttribDict
from blixt_rp.core.log_curve_new import handle_depth
from blixt_utils.utils import print_info, add_one, fix_well_name, cycle_colors


class StratUnit(object):
    """
    Contains information about one specific stratigraphic unit (Group, Formation, Member, ...)
    Must have a unique name
    """

    def __init__(self,
                 name: str,
                 level: int,
                 desc=None,
                 source=None,
                 color=None
                 ):
        """

        :param name:
        :param level:
            int
            High number indicates that high up in the hierarchy (e.g. "Era" or "Group") while lower (even negative)
            numbers indicate that is a sub interval (e.g. "Stage" or "Member")
            Set it to None when unknown or uncertain
        :param desc:
        :param source:
            str
            E.G. <User name>, or  "sodir", ...
        :param color:
        """
        if name is None:
            raise IOError('An Interval must have a name')
        self.name = name
        self.level = level
        self.desc = desc
        self.source = source
        self.color = color

    def __str__(self):
        return print_function(self)

    def __getitem__(self, item):
        return self.__dict__[item]

    def keys(self):
        return self.__dict__.keys()

class Interval:
    """
   Creates a relation between an interval and a well, and the top and base in that well
    """
    def __init__(self,
                 well: str,
                 top: pint.Quantity,
                 base: pint.Quantity,
                 interval_info: StratUnit,
                 depth_type: str | None = None):
        """

        :param well:
        :param top:
        :param base:
        :param interval_info:
        :param depth_type:
            str
        """
        self.well = well
        if depth_type is None:
            depth_type = 'md'
        self.top = handle_depth(top, depth_units=None, depth_type=depth_type)
        self.base = handle_depth(base, depth_units=None, depth_type=depth_type)
        self.interval_info = interval_info

    def __str__(self):
        return print_function(self)
    @property
    def name(self):
        return self.interval_info.name

    # Dont know if a need setters?
    # @name.setter
    # def name(self, value):
    #     self.interval_info.name = value

    @property
    def level(self):
        return self.interval_info.level

    @property
    def color(self):
        return self.interval_info.color

    @property
    def source(self):
        return self.interval_info.source

    @property
    def desc(self):
        return self.interval_info.desc

    @property
    def thickness(self):
        return self.base - self.top

    @property
    def mid(self):
        return 0.5 * (self.top + self.base)

    def distance_to_top(self, this_depth: pint.Quantity):
        """
        Distance to top of interval.
        Positive values when this_depth is deeper than top
        :param this_depth:
            pint.Quantity
        :return:
            pint.Quantity
        """
        return -1. * (self.top - this_depth)  # putting self.top ensures the result uses the same units as self.top

    def distance_to_base(self, this_depth: pint.Quantity):
        """
        Distance to base of interval.
        Positive values when this_depth is shallower than base
        :param this_depth:
            pint.Quantity
        :return:
            pint.Quantity
        """
        return self.base - this_depth

class Intervals(object):
    """
    Object that holds a number of working intervals / tops for many wells
    """
    def __init__(self,
                 name: str | None = None,
                 intervals: list | None = None,
                 verbose=False):
        """
        :param name:
            str
        :param intervals:
            list list(Interval)
            List of all single intervals
        :param verbose:
            bool
        """
        self.name = name
        if intervals is None:
            self.intervals = []
        elif isinstance(intervals, list):
            self.intervals = intervals
        else:
            raise IOError('Intervals must be a list (of intervals)')

    def __len__(self):
        return len(self.intervals)

    def __str__(self):
        return print_function(self)

    def interval_names(self):
        return list(set([_x.name for _x in self.intervals]))

    def well_names(self):
        return list(set([_x.well for _x in self.intervals]))

    def get_interval(self, interval_name, well_name):
        for _interval in self.intervals:
            if _interval.name == interval_name and _interval.well == well_name:
                return _interval
        return None

    def get_well_depth_range(self,  well_name):
        min_top = 1E6
        max_base = -1E6
        for _interval in self.intervals:
            if _interval.well != well_name:
                continue
            if _interval.top.magnitude <= min_top:
                min_top = _interval.top.magnitude
            if _interval.base.magnitude >= max_base:
                max_base = _interval.base.magnitude
        return min_top,  max_base

    def get_well_intervals(self, well_name) -> list:
        _levels = []
        for _interval in self.intervals:
            if _interval.well == well_name:
                _levels.append(_interval)
        return _levels

    def add_interval(self, interval: Interval):
        self.intervals.append(interval)

    def get_info_dict(self) -> dict:
        """
        Returns a dictionary which contains all the StratUnit content
        :return:
        """
        interval_info_dict = {'name': [], 'level': [], 'desc': [], 'source': [], 'color': []}
        for _interval in self.intervals:
            if _interval.name in interval_info_dict['name']:
                continue
            else:
                interval_info_dict['name'].append(_interval.name)
                interval_info_dict['level'].append(_interval.level)
                interval_info_dict['desc'].append(_interval.desc)
                interval_info_dict['source'].append(_interval.source)
                interval_info_dict['color'].append(_interval.color)
        return interval_info_dict

    def get_intervals_dict(self,
                           well_name: str | None = None) -> dict:
        """
        Returns a dictionary with all intervals
        if well_name is not None, it filters out all other wells
        :return:
        """
        intervals_dict = {'well': [], 'name': [], 'top MD [m]': [], 'base MD [m]': [],
                          'level': [], 'source': [], 'note': []}
        for _interval in self.intervals:
            if well_name is not None:
                if well_name != _interval.well:
                    continue
            intervals_dict['well'].append(_interval.well)
            intervals_dict['name'].append(_interval.name)
            intervals_dict['top MD [m]'].append(_interval.top.to('m').magnitude)
            intervals_dict['base MD [m]'].append(_interval.base.to('m').magnitude)
            intervals_dict['level'].append(_interval.level)
            intervals_dict['source'].append(_interval.source)
            intervals_dict['note'].append(_interval.desc)
        return intervals_dict

    def write_to_excel(self, file_name, intervals_sheet, interval_info_sheet, append: bool = False):
        import openpyxl
        int_rows = 0
        int_info_rows = 0
        int_header = True
        info_header = True
        if not os.path.isfile(file_name):
            wb = openpyxl.Workbook()
            wb.save(filename=file_name)
        elif append:
            # Figure out the row numbers where we can start appending data
            wb = openpyxl.load_workbook(file_name)
            try:
                sheet = wb[intervals_sheet]
                int_rows = sheet.max_row
                int_header = False
            except KeyError:   # this sheet doesn't exist, so we create it later
                pass
            try:
                sheet = wb[interval_info_sheet]
                int_info_rows = sheet.max_row
                info_header = False
            except KeyError:   # this sheet doesn't exist, so we create it later
                pass
            wb.close()

        # First collect and write the interval info
        interval_info_dict = self.get_info_dict()
        df = pd.DataFrame(interval_info_dict)
        with pd.ExcelWriter(file_name, mode='a', if_sheet_exists='overlay', engine='openpyxl') as writer:
            df.to_excel(
                writer, sheet_name=interval_info_sheet, startcol=0, startrow=int_info_rows, index=False,
                header=info_header)

        # Then the individual intervals in each well
        intervals_dict = self.get_intervals_dict()
        df = pd.DataFrame(intervals_dict)
        with pd.ExcelWriter(file_name, mode='a', if_sheet_exists='overlay', engine='openpyxl') as writer:
            df.to_excel(
                writer, sheet_name=intervals_sheet, startcol=0, startrow=int_rows, index=False,
                header=int_header)

    def read_sodir_tops(self, file_name, testing=True):
        df = pd.read_excel(file_name, engine='openpyxl')
        for i, well_name in enumerate(df['Wellbore name']):
            well_name = fix_well_name(well_name)
            _top = df['Top depth [m]'][i]
            _base = df['Bottom depth [m]'][i]
            _name = df['Lithostrat. unit'][i]
            _level = df['Level'][i]
            if _level == 'GROUP':
                l = 3
            elif _level == 'FORMATION':
                l = 2
            else:
                l = 0

            _interval_info = StratUnit(_name, l, source='SoDir')
            _interval = Interval(well_name, Q_(float(_top), 'm'), Q_(float(_base), 'm'),
                                       interval_info=_interval_info)

            self.add_interval(_interval)

            if testing and (i > 40):
                break

    def read_blixt_tops(self, file_name: str, intervals_sheet: str = 'Working intervals',
                        interval_info_sheet: str = 'Interval info'):
        # see if there is a sheet with interval info
        interval_info = False
        if interval_info_sheet in list(pd.read_excel(file_name, engine='openpyxl', sheet_name=None)):
            raise NotImplementedError('Using the interval info sheet has not been taken into use yet')
        df = pd.read_excel(file_name, engine='openpyxl', sheet_name=intervals_sheet, header=4)
        for i, well_name in enumerate(df['Given well name']):
            interval_name = df['Interval name'][i]
            if not interval_info:
                this_level = get_level_from_name(interval_name, source='sodir')
                interval_info = StratUnit(interval_name, this_level, source='blixt')
            _interval = Interval(
                well=well_name,
                top=Q_(float(df['Top depth'][i]), 'm'),
                base=Q_(float(df['Base depth'][i]), 'm'),
                interval_info=interval_info
            )

            self.add_interval(_interval)

    def bokeh_plot(self,
                   well_name: str,
                   p: figure):


def print_function(my_object):
    keys = list(my_object.__dict__.keys())
    try:
        i = max([len(k) for k in keys])
    except ValueError:
        # no keys
        return ''
    pattern = "%%%ds: %%s" % (i)
    head = [pattern % (k, str(my_object.__dict__[k])) for k in keys]
    return "\n".join(head)

def get_level_from_name(_name: str, source: str | None = None) -> int:
    if source is None:
        source = 'sodir'  # Norwegian "Sokkel direktoratet"
    this_level = 0
    if source == 'sodir':
        # Try to extract level based on name of interval
        if ' gp' in _name.lower():
            this_level = 2
        elif ' fm' in _name.lower():
            this_level = 1
    return this_level


